import sys
from PyQt5.QtWidgets import QApplication, QMainWindow, QTextEdit, QPushButton, QVBoxLayout, QWidget, QFileDialog
from docx import Document
from docx.shared import RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

class TableConverter(QMainWindow):
    def __init__(self):
        super().__init__()
        self.initUI()

    def initUI(self):
        self.setWindowTitle('表格转换工具')
        self.setGeometry(100, 100, 800, 600)

        # 创建中央部件和布局
        central_widget = QWidget()
        self.setCentralWidget(central_widget)
        layout = QVBoxLayout(central_widget)

        # 创建文本编辑框
        self.text_edit = QTextEdit()
        self.text_edit.setPlaceholderText('请在此输入Markdown表格...')
        layout.addWidget(self.text_edit)

        # 创建保存按钮
        save_button = QPushButton('保存为Word文档')
        save_button.clicked.connect(self.save_to_word)
        layout.addWidget(save_button)

    def remove_markdown_bold(self, text):
        """移除Markdown加粗语法"""
        while '**' in text:
            text = text.replace('**', '')
        return text

    def parse_markdown_tables(self, text):
        """解析Markdown文本中的表格"""
        tables = []
        current_table = []
        
        for line in text.split('\n'):
            line = line.strip()
            if line.startswith('|'):
                # 检查是否为分隔行（包含 - 的行）
                if not any(cell.strip().replace('-', '').replace(':', '') == '' for cell in line.strip('|').split('|')):
                    # 移除首尾的|并分割，同时处理加粗语法
                    row = [self.remove_markdown_bold(cell.strip()) for cell in line.strip('|').split('|')]
                    current_table.append(row)
            elif current_table:
                if len(current_table) > 0:  # 修改条件，只要有内容就添加
                    tables.append(current_table)
                current_table = []

        if current_table and len(current_table) > 0:
            tables.append(current_table)

        return tables

    def create_word_table(self, doc, table_data):
        """创建Word表格并设置格式"""
        if not table_data:
            return

        # 创建表格
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.style = 'Table Grid'

        # 填充数据并设置格式
        for i, row in enumerate(table_data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell
                paragraph = table.cell(i, j).paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # 设置表头背景色为浅灰色
                if i == 0:
                    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="D3D3D3"/>')
                    table.cell(i, j)._tc.get_or_add_tcPr().append(shading_elm)

        doc.add_paragraph()  # 添加空行分隔表格

    def save_to_word(self):
        """保存为Word文档"""
        markdown_text = self.text_edit.toPlainText()
        if not markdown_text:
            return

        # 解析表格数据
        tables = self.parse_markdown_tables(markdown_text)
        
        # 创建保存文件对话框
        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "保存Word文档",
            "",
            "Word Documents (*.docx)"
        )

        if file_path:
            # 创建Word文档
            doc = Document()
            
            # 添加所有表格
            for table_data in tables:
                self.create_word_table(doc, table_data)

            # 保存文档
            try:
                doc.save(file_path)
                self.statusBar().showMessage('文档保存成功！', 3000)
            except Exception as e:
                self.statusBar().showMessage(f'保存失败：{str(e)}', 3000)

def main():
    app = QApplication(sys.argv)
    ex = TableConverter()
    ex.show()
    sys.exit(app.exec_())

if __name__ == '__main__':
    main()
